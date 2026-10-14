"""Exercise the Excel review workflow from source export through final parsing."""
from __future__ import annotations

import contextlib
import io
import json
import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
for directory in (BASE / "leave", BASE / "finally"):
    sys.path.insert(0, str(directory))
import calc_leave
import calc_finally
from cross_month_leave_test import schedule_context


class LeaveReviewWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "请假明细表.xlsx"
        self.ctx = schedule_context(2026, 9, {date(2026, 9, day) for day in (1, 2, 3, 4)})

    def row(self, hours=8, *, approval="A1", kind="事假", start=None, end=None, status="完成"):
        return (
            "MT0001", "测试员工", "总部", "产品部", "研发组", kind,
            start or datetime(2026, 9, 1, 9), end or datetime(2026, 9, 1, 18, 30),
            hours, datetime(2026, 8, 25, 9), datetime(2026, 8, 26, 9),
            approval, status, "同意", False, 2,
        )

    def generate(self, rows):
        header = (
            "发起人工号", "发起人姓名", "一级部门", "二级部门", "三级部门", "请假类型",
            "开始时间", "结束时间", "系统时长", "发起时间", "完成时间", "审批编号",
            "审批状态", "审批结果", "是否实习生", "源文件行号",
        )
        with contextlib.redirect_stdout(io.StringIO()):
            calc_leave.process([header] + rows, str(self.path), self.ctx)

    def edit(self, changes, row=2):
        wb = openpyxl.load_workbook(self.path)
        ws = wb["请假明细"]
        headers = {cell.value: cell.column for cell in ws[1]}
        for key, value in changes.items():
            ws.cell(row, headers[key], value)
        wb.save(self.path)
        wb.close()

    def values(self):
        wb = openpyxl.load_workbook(self.path, data_only=True)
        ws = wb["请假明细"]
        data = dict(zip((cell.value for cell in ws[1]), (cell.value for cell in ws[2])))
        count = wb["异常清单"].max_row - 1
        wb.close()
        return data, count

    def test_matching_same_month_is_automatic_and_ledger_is_shared(self):
        self.generate([self.row()])
        values, issues = self.values()
        self.assertEqual(values["复核状态"], "自动通过")
        self.assertEqual(issues, 0)
        self.assertEqual(calc_finally.parse_leave_summary(str(self.path), self.ctx)["MT0001"]["事假"], 1)
        details = calc_finally.parse_leave_day_details(str(self.path), self.ctx)
        self.assertEqual(details["MT0001"][date(2026, 9, 1)]["事假"], 1)

    def test_mismatch_is_blank_and_final_stops(self):
        self.generate([self.row(4)])
        values, issues = self.values()
        self.assertEqual(values["系统时长"], 4)
        self.assertEqual(values["作息时长"], 8)
        self.assertIsNone(values["最终请假时长"])
        self.assertEqual(issues, 1)
        with self.assertRaisesRegex(ValueError, "尚未确认"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_confirmation_changes_both_summary_and_daily_detail(self):
        self.generate([self.row(4)])
        self.edit({"复核状态": "已确认", "确认请假时长": 4, "确认说明": "核对审批，实际请半天"})
        summary = calc_finally.parse_leave_summary(str(self.path), self.ctx)
        detail = calc_finally.parse_leave_day_details(str(self.path), self.ctx)
        self.assertEqual(summary["MT0001"]["事假"], 0.5)
        self.assertEqual(detail["MT0001"][date(2026, 9, 1)]["事假"], 0.5)

    def test_zero_confirmation_does_not_fall_back_to_system(self):
        self.generate([self.row(4)])
        self.edit({"复核状态": "已确认", "确认请假时长": 0, "确认说明": "重复申请，本月不入账"})
        self.assertEqual(calc_finally.parse_leave_summary(str(self.path), self.ctx), {})
        self.assertEqual(calc_finally.parse_leave_day_details(str(self.path), self.ctx), {})

    def test_state_only_and_missing_reason_cannot_bypass_review(self):
        self.generate([self.row(4)])
        self.edit({"复核状态": "已确认"})
        with self.assertRaisesRegex(ValueError, "确认说明"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)
        self.edit({"确认说明": "已核对"})
        with self.assertRaisesRegex(ValueError, "确认值"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_bad_system_values_are_visible_pending_records(self):
        for amount in (None, "", "坏数据", "3..5小时", "3..5天", -1, "NaN", "inf", 80):
            with self.subTest(amount=amount):
                self.generate([self.row(amount)])
                values, _ = self.values()
                self.assertEqual(values["复核状态"], "待确认")
                self.assertIsNone(values["最终请假天数"])

    def test_missing_ids_do_not_collapse_multiple_source_rows(self):
        path = Path(self.tmp.name) / "source.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["审批状态", "审批结果", "发起人工号", "发起人姓名", "发起人部门", "请假类型", "开始时间", "结束时间", "时长", "审批编号"])
        for person in ("MT1", "MT2"):
            ws.append(["完成", "同意", person, person, "总部", "事假", "2026-09-01", "2026-09-01", "8小时", None])
        wb.save(path)
        wb.close()
        with contextlib.redirect_stdout(io.StringIO()):
            rows = calc_leave.clean_export(str(path))
        self.assertEqual(len(rows), 3)

    def test_latest_revocation_invalidates_old_approved_version(self):
        path = Path(self.tmp.name) / "source.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["审批状态", "审批结果", "发起人工号", "发起人姓名", "发起人部门", "请假类型", "开始时间", "结束时间", "时长", "审批编号", "完成时间"])
        for status, finish in (("已撤销", "2026-09-03 12:00"), ("完成", "2026-09-02 12:00")):
            ws.append([status, "同意", "MT1", "员工", "总部", "事假", "2026-09-01", "2026-09-01", "8小时", "A1", finish])
        wb.save(path)
        wb.close()
        with contextlib.redirect_stdout(io.StringIO()):
            rows = calc_leave.clean_export(str(path))
        self.assertEqual(len(rows), 1)

    def test_pending_approval_and_modified_status_require_review(self):
        for status in ("审批中", "已修改", "未知状态"):
            self.generate([self.row(status=status)])
            self.assertEqual(self.values()[0]["复核状态"], "待确认")

    def test_different_overlapping_approvals_are_not_merged(self):
        self.generate([self.row(approval="A1"), self.row(approval="A2")])
        self.assertEqual(self.values()[1], 2)
        self.edit({"复核状态": "已确认", "确认请假时长": 8, "确认说明": "保留有效申请"})
        self.edit({"复核状态": "已确认", "确认请假时长": 0, "确认说明": "重复申请"}, row=3)
        self.assertEqual(calc_finally.parse_leave_summary(str(self.path), self.ctx)["MT0001"]["事假"], 1)

    def test_natural_leave_preserves_days_without_fake_hours(self):
        self.generate([self.row("3天", kind="陪产假", end=datetime(2026, 9, 3, 18, 30))])
        values, _ = self.values()
        self.assertEqual(values["计量口径"], "自然日")
        self.assertIsNone(values["最终请假时长"])
        self.assertEqual(values["最终请假天数"], 3)
        self.assertEqual(calc_finally.parse_leave_summary(str(self.path), self.ctx)["MT0001"]["陪产假"], 3)

    def test_unknown_day_rule_requires_review(self):
        self.generate([self.row("3天", kind="婚假")])
        self.assertEqual(self.values()[0]["复核状态"], "待确认")

    def test_multiday_correction_requires_conserving_daily_ledger(self):
        self.generate([self.row(12, end=datetime(2026, 9, 2, 18, 30))])
        self.edit({"复核状态": "已确认", "确认请假时长": 12, "确认说明": "首日4小时次日8小时"})
        with self.assertRaisesRegex(ValueError, "确认日分摊"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)
        self.edit({"确认日分摊": "2026-09-01=4;2026-09-02=8"})
        details = calc_finally.parse_leave_day_details(str(self.path), self.ctx)["MT0001"]
        self.assertEqual(details[date(2026, 9, 1)]["事假"], 0.5)
        self.assertEqual(details[date(2026, 9, 2)]["事假"], 1)

    def test_cross_month_does_not_round_system_total_up(self):
        start, end = datetime(2026, 8, 31, 9), datetime(2026, 9, 1, 12)
        allocation, _ = calc_leave.allocate_adjacent_cross_month_system_daily_hours(start, end, 11.2)
        self.assertAlmostEqual(sum(allocation.values()), 11.2)
        self.generate([self.row(11.2, start=start, end=end)])
        self.assertEqual(self.values()[0]["最终请假时长"], 3.2)

    def test_three_month_allocation_conserves_total(self):
        start, end = datetime(2026, 7, 31, 9), datetime(2026, 9, 1, 18, 30)
        total = sum(8 for offset in range((end.date() - start.date()).days + 1) if (start.date() + timedelta(days=offset)).weekday() < 5)
        allocation, _ = calc_leave.allocate_adjacent_cross_month_system_daily_hours(start, end, total)
        self.assertEqual(sum(allocation.values()), total)
        self.assertEqual({(day.year, day.month) for day in allocation}, {(2026, 7), (2026, 8), (2026, 9)})

    def test_invalid_dates_stay_in_exception_list(self):
        self.generate([self.row(start="坏日期")])
        self.assertEqual(self.values()[1], 1)
        self.generate([self.row(start=datetime(2026, 9, 2), end=datetime(2026, 9, 1))])
        self.assertEqual(self.values()[0]["复核状态"], "待确认")

    def test_wrong_month_and_tampered_ledger_are_rejected(self):
        self.generate([self.row()])
        with self.assertRaisesRegex(ValueError, "月份"):
            calc_finally.parse_leave_summary(str(self.path), schedule_context(2026, 10, set()))
        self.edit({"日分摊": json.dumps({"2026-10-01": {"hours": 8, "days": 1}})})
        with self.assertRaisesRegex(ValueError, "日分摊"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_cross_month_confirmation_requires_whole_segment_total(self):
        self.generate([self.row(None, start=datetime(2026, 8, 31, 9))])
        self.edit({"复核状态": "已确认", "确认请假时长": 8, "确认说明": "核对源审批实际整段16小时"})
        with self.assertRaisesRegex(ValueError, "整段"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)
        self.edit({"确认整段时长": 16, "确认日分摊": "2026-08-31=8;2026-09-01=8"})
        self.assertEqual(calc_finally.parse_leave_summary(str(self.path), self.ctx)["MT0001"]["事假"], 1)
        self.edit({"确认整段时长": 20})
        with self.assertRaisesRegex(ValueError, "整段"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_confirmation_does_not_hide_out_of_range_daily_allocation(self):
        self.generate([self.row(4)])
        self.edit({"复核状态": "已确认", "确认请假时长": 4, "确认说明": "核对", "确认日分摊": "2026-09-02=4"})
        with self.assertRaisesRegex(ValueError, "超出"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_cross_month_schedule_disagreement_is_pending(self):
        self.ctx["main_working_days"].remove(date(2026, 9, 2))
        self.generate([self.row(32, start=datetime(2026, 8, 31, 9), end=datetime(2026, 9, 3, 18, 30))])
        self.assertEqual(self.values()[0]["复核状态"], "待确认")

    def test_legacy_sheet_missing_review_fields_is_rejected(self):
        wb = openpyxl.Workbook()
        wb.active.append(["工号", "姓名", "请假类型", "最终请假天数"])
        wb.active.append(["MT0001", "员工", "事假", 1])
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "复核状态"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_duplicate_approval_with_different_amounts_is_rejected(self):
        self.generate([self.row(approval="A1"), self.row(approval="A1")])
        self.edit({"复核状态": "已确认", "确认请假时长": 8, "确认说明": "核对有效版本"})
        self.edit({"复核状态": "已确认", "确认请假时长": 4, "确认说明": "冲突版本"}, row=3)
        with self.assertRaisesRegex(ValueError, "同一审批编号"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_confirmed_overlaps_cannot_exceed_one_day(self):
        self.generate([self.row(approval="A1"), self.row(approval="A2")])
        for row in (2, 3):
            self.edit({"复核状态": "已确认", "确认请假时长": 8, "确认说明": "已核对"}, row=row)
        with self.assertRaisesRegex(ValueError, "超过1天"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_fractional_approvals_share_precision_with_day_ledger(self):
        self.generate([
            self.row(3, approval="A1", end=datetime(2026, 9, 1, 12)),
            self.row(3, approval="A2", start=datetime(2026, 9, 1, 14), end=datetime(2026, 9, 1, 17)),
        ])
        summary = calc_finally.parse_leave_summary(str(self.path), self.ctx)
        daily = calc_finally.parse_leave_day_details(str(self.path), self.ctx)
        self.assertEqual(summary["MT0001"]["事假"], 0.75)
        self.assertEqual(daily["MT0001"][date(2026, 9, 1)]["事假"], 0.75)

    def test_partial_maternity_preserves_fractional_attendance_days(self):
        self.generate([self.row("1天", kind="产假", end=datetime(2026, 9, 1, 12))])
        summary = calc_finally.parse_leave_summary(str(self.path), self.ctx)
        daily = calc_finally.parse_leave_day_details(str(self.path), self.ctx)
        self.assertEqual(summary["MT0001"]["产假"], 0.375)
        self.assertEqual(daily["MT0001"][date(2026, 9, 1)]["产假"], 0.375)

    def test_unrecognized_workbook_cannot_be_treated_as_no_leave(self):
        wb = openpyxl.Workbook()
        wb.active.append(["无效表头"])
        wb.active.append(["无效数据"])
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "未找到"):
            calc_finally.parse_leave_summary(str(self.path), self.ctx)

    def test_modified_version_without_completion_time_cannot_revive_old_approval(self):
        path = Path(self.tmp.name) / "source.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["审批状态", "审批结果", "发起人工号", "发起人姓名", "发起人部门", "请假类型", "开始时间", "结束时间", "时长", "审批编号", "完成时间"])
        for status, finish in (("完成", "2026-09-02 12:00"), ("已修改", None)):
            ws.append([status, "同意", "MT1", "员工", "总部", "事假", "2026-09-01", "2026-09-01", "8小时", "A1", finish])
        wb.save(path)
        wb.close()
        with contextlib.redirect_stdout(io.StringIO()):
            rows = calc_leave.clean_export(str(path))
            calc_leave.process(rows, str(self.path), self.ctx)
        self.assertEqual(self.values()[0]["复核状态"], "待确认")
        self.assertIn("版本待确认", self.values()[0]["异常原因"])


if __name__ == "__main__":
    unittest.main()

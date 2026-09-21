# -*- coding: utf-8 -*-
"""Tests for the part-time monthly punch renderer (req: 生成的 Excel 可以正常打开，并包含关键表头和目标月份数据)."""
from __future__ import annotations

import sys
import shutil
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent
PARTTIME_ROOT = ROOT / "parttime"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(PARTTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(PARTTIME_ROOT))

import calc_parttime_summary as parttime  # noqa: E402
import parttime_monthly_punch as punch  # noqa: E402
import runner  # noqa: E402


@contextmanager
def temporary_workdir():
    workdir = ROOT / f".parttime-test-{uuid.uuid4().hex}"
    workdir.mkdir()
    try:
        yield workdir
    finally:
        shutil.rmtree(workdir)


def write_and_parse(cfg: dict) -> dict:
    with temporary_workdir() as workdir:
        punch.render(workdir, cfg)
        out = workdir / "outputs" / f"兼职月度打卡记录_{cfg['year']}{cfg['month']:02d}.xlsx"
        return parttime.parse_attendance_detail(str(out))


class ParttimeMonthlyPunchRenderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.cfg = {
            "year": 2026,
            "month": 7,
            "days_in_month": 31,
            "matched": [
                {
                    "name": "张三",
                    "employee_no": "MT001",
                    "position": "兼职",
                    "department": "北京地铁",
                    "matched_by": "employee_no",
                    "days": {1: "正常 (08:30,18:00)", 2: "迟到 (08:35,18:00)", 15: "旷工"},
                },
                {
                    "name": "李四",
                    "employee_no": "",
                    "position": "",
                    "department": "",
                    "matched_by": "name",
                    "days": {1: "正常 (09:00,17:30)"},
                },
            ],
            "unmatched": [
                {"name": "王五", "employee_no": "MT003", "position": "兼职", "department": ""},
            ],
            "anomalies": ["姓名「赵六」在打卡记录中出现多次且工号缺失"],
        }

    def test_render_produces_readable_workbook(self) -> None:
        with temporary_workdir() as workdir:
            result = punch.render(workdir, self.cfg)
            self.assertTrue(Path(result["path"]).exists())
            self.assertTrue(result["file_name"].endswith(".xlsx"))

    def test_parser_consumes_grid_with_key_headers(self) -> None:
        parsed = write_and_parse(self.cfg)
        # 关键表头对应的数据行应被解析（姓名/工号列存在）。
        self.assertIn("张三", parsed)
        self.assertIn("李四", parsed)

    def test_target_month_data_populated(self) -> None:
        parsed = write_and_parse(self.cfg)
        self.assertEqual(parttime._entry_value(parsed["张三"].get(1)), 1.0)
        self.assertEqual(parttime._entry_value(parsed["张三"].get(2)), 1.0)
        # 旷工日记为 0 出勤。
        self.assertEqual(parttime._entry_value(parsed["张三"].get(15)), 0.0)

    def test_parser_accepts_safe_aliases_and_locates_day_columns(self) -> None:
        with temporary_workdir() as workdir:
            path = workdir / "考勤明细_别名.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "考勤明细"
            ws.append(["考勤时间：2026-07-01 至 2026-07-31"])
            ws.append(["员工姓名", "员工编号", "部门名称", "岗位名称", "1", "07-02", "15", "29"])
            ws.append([None, None, None, None, None, None, None, None])
            ws.append([
                "实习生甲",
                "JZ001",
                "研发部",
                "实习生",
                "正常 (09:00,18:00)",
                "旷工",
                "正常 (09:00,18:00)",
                "正常 (09:00,18:00)",
            ])
            wb.save(path)
            wb.close()

            parsed = parttime.parse_attendance_detail(str(path))
            validation = runner._check_parttime_detail_headers(str(path))

        self.assertIn("实习生甲", parsed)
        self.assertEqual(parttime._entry_value(parsed["实习生甲"].get(1)), 1.0)
        self.assertEqual(parttime._entry_value(parsed["实习生甲"].get(2)), 0.0)
        self.assertEqual(parttime._entry_value(parsed["实习生甲"].get(15)), 1.0)
        self.assertEqual(parttime._entry_value(parsed["实习生甲"].get(29)), 1.0)
        self.assertTrue(validation["ok"], validation)

    def test_parser_accepts_simple_name_day_matrix_without_employee_code(self) -> None:
        with temporary_workdir() as workdir:
            path = workdir / "员工考勤表.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "考勤表"
            ws.append(["姓名", "1", "15", "29", "出勤天数", "所属组织"])
            ws.append(["员工甲", "√", "√", "√", 3, "南京地铁"])
            ws.append(["员工乙", 1, None, 1, 2, "南京地铁"])
            wb.save(path)
            wb.close()

            parsed = parttime.parse_attendance_detail(str(path))
            validation = runner._check_parttime_detail_headers(str(path))

        self.assertEqual(parttime._entry_value(parsed["员工甲"].get(1)), 1.0)
        self.assertEqual(parttime._entry_value(parsed["员工甲"].get(15)), 1.0)
        self.assertEqual(parttime._entry_value(parsed["员工甲"].get(29)), 1.0)
        self.assertEqual(parttime._entry_value(parsed["员工乙"].get(29)), 1.0)
        self.assertTrue(validation["ok"], validation)

    def test_parser_rejects_wide_name_only_sheet_without_day_headers(self) -> None:
        with temporary_workdir() as workdir:
            path = workdir / "非考勤宽表.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["姓名", "字段1", "字段2", "字段3", "字段4", "字段5", "字段6"])
            ws.append(["员工甲", 1, 2, 3, 4, 5, 6])
            wb.save(path)
            wb.close()

            validation = runner._check_parttime_detail_headers(str(path))

        self.assertFalse(validation["ok"], validation)

    def test_unmatched_default_position_keeps_row_in_scope(self) -> None:
        # 空 职位 应默认填充为「兼职」，使行能通过兼职汇总的范围过滤。
        parsed = write_and_parse(self.cfg)
        self.assertIn("李四", parsed)

    def test_empty_matched_still_renders(self) -> None:
        cfg = dict(self.cfg, matched=[], unmatched=self.cfg["unmatched"], anomalies=[])
        parsed = write_and_parse(cfg)
        self.assertEqual(parsed, {})

    def test_invalid_month_rejected(self) -> None:
        with self.assertRaises(ValueError):
            punch.render(Path(__file__).resolve().parent, dict(self.cfg, month=13))


if __name__ == "__main__":
    unittest.main()

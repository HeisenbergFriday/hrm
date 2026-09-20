# -*- coding: utf-8 -*-
"""Regression tests for part-time attendance priority rules."""
from __future__ import annotations

import sys
import re
import tempfile
import unittest
import zipfile
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel


ROOT = Path(__file__).resolve().parent
PARTTIME_ROOT = ROOT / "parttime"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(PARTTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(PARTTIME_ROOT))

import calc_parttime_summary as parttime  # noqa: E402


class OffsitePersonalLeavePriorityTests(unittest.TestCase):
    def test_business_trip_with_personal_leave_is_not_counted(self) -> None:
        text = "出差07-23 09:00到07-23 18:30 1.0天,事假07-23 09:00到07-23 18:30 8小时\n(-)"

        result = parttime._parse_daily_text_value(
            text,
            count_outing_as_present=True,
        )

        self.assertIsNone(result)

    def test_outing_with_personal_leave_is_not_counted(self) -> None:
        text = "外出07-23 09:00到07-23 18:30,事假07-23 09:00到07-23 18:30 8小时\n(-)"

        result = parttime._parse_daily_text_value(
            text,
            count_outing_as_present=True,
        )

        self.assertIsNone(result)

    def test_offsite_without_personal_leave_still_counts_as_present(self) -> None:
        for text in ("外出07-23 09:00到07-23 18:30\n(-)", "出差07-23 09:00到07-23 18:30 1.0天\n(-)"):
            with self.subTest(text=text):
                result = parttime._parse_daily_text_value(
                    text,
                    count_outing_as_present=True,
                )

                self.assertIsNotNone(result)
                self.assertEqual(parttime._entry_value(result), 1.0)


class FixedScheduleAttendanceTests(unittest.TestCase):
    def test_default_fixed_names_excludes_chang_yufan(self) -> None:
        self.assertEqual(
            parttime.DEFAULT_SPECIAL_DEFAULT_NAMES,
            ("王心英", "刘芮", "汤颖", "周代林", "陈富庆"),
        )
        self.assertEqual(parttime.DEFAULT_SPECIAL_CHENGDU_NAMES, set())

    def test_formula_schedule_uses_cached_dates_for_full_month(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "8月作息表.xlsx"
            self._write_formula_schedule(path)

            day_maps = parttime._build_default_day_maps(2026, 8, str(path))

        self.assertEqual(len(day_maps["main"]), 24)
        self.assertEqual(parttime._entry_value(day_maps["main"].get(15)), 1.0)
        self.assertEqual(parttime._entry_value(day_maps["main"].get(29)), 1.0)

    def test_fixed_person_keeps_schedule_value_and_only_reminds_anomalies(self) -> None:
        monthly_summary = {
            "刘芮": {
                15: parttime._parse_daily_text_value("标准:上班迟到12分钟\n(09:12,18:30)"),
                29: parttime._parse_daily_text_value("标准:旷工\n(-)"),
            }
        }
        schedule_days = {
            15: parttime._make_daily_entry(1.0),
            29: parttime._make_daily_entry(1.0),
        }

        resolved, source, warnings = parttime._resolve_row_days(
            {"姓名": "刘芮", "所属公司": ""},
            attendance_detail={},
            monthly_summary=monthly_summary,
            schedules={},
            default_day_maps={"main": schedule_days, "chengdu": schedule_days},
            special_default_name_set={"刘芮"},
        )

        self.assertEqual(source, parttime.SOURCE_SPECIAL_DEFAULT_WEEKDAYS)
        self.assertEqual(warnings, [])
        self.assertEqual(parttime._entry_value(resolved[15]), 1.0)
        self.assertEqual(parttime._entry_value(resolved[29]), 1.0)
        self.assertEqual(sum(parttime._entry_value(entry) or 0 for entry in resolved.values()), 2.0)
        self.assertIn("迟到12分钟", parttime._entry_note(resolved[15]))
        self.assertIn("旷工", parttime._entry_note(resolved[29]))
        self.assertNotIn("扣减", parttime._entry_note(resolved[15]))
        self.assertNotIn("未计出勤", parttime._entry_note(resolved[29]))
        self.assertTrue(parttime._is_alert_entry(resolved[15]))
        self.assertFalse(parttime._is_deducted_entry(resolved[15]))
        self.assertFalse(parttime._is_deducted_entry(resolved[29]))

    @staticmethod
    def _write_formula_schedule(path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "8月作息表"
        ws.cell(1, 1, "2026年8月作息时间表")
        for col_idx, header in enumerate(("周数", "周一", "周二", "周三", "周四", "周五", "周六", "周日"), start=1):
            ws.cell(2, col_idx, header)

        yellow = PatternFill("solid", fgColor="FFFF00")
        start = date(2026, 7, 27)
        cached_values: dict[str, int] = {}
        for offset in range(42):
            current = start + timedelta(days=offset)
            row_idx = 3 + offset // 7
            col_idx = 2 + offset % 7
            if offset % 7 == 0:
                ws.cell(row_idx, 1, offset // 7 + 1)
            cell = ws.cell(row_idx, col_idx)
            if offset < 14:
                cell.value = current
            else:
                previous_col = col_idx - 1 if col_idx > 2 else 8
                previous_row = row_idx if col_idx > 2 else row_idx - 1
                cell.value = f"={get_column_letter(previous_col)}{previous_row}+1"
                cached_values[cell.coordinate] = int(to_excel(current))
            cell.number_format = "d"
            if current.month == 8 and (current.weekday() < 5 or current.day in {1, 15, 29}):
                cell.fill = yellow

        wb.save(path)
        wb.close()

        patched_path = path.with_suffix(".cached.xlsx")
        with zipfile.ZipFile(path, "r") as source_zip, zipfile.ZipFile(patched_path, "w") as target_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    xml = data.decode("utf-8")
                    for coordinate, cached_value in cached_values.items():
                        pattern = rf'(<c r="{coordinate}"[^>]*>.*?<f>.*?</f><v>).*?(</v>)'
                        xml, count = re.subn(
                            pattern,
                            rf"\g<1>{cached_value}\g<2>",
                            xml,
                            count=1,
                        )
                        if count != 1:
                            raise AssertionError(f"未能写入公式缓存值: {coordinate}")
                    data = xml.encode("utf-8")
                target_zip.writestr(info, data)
        patched_path.replace(path)


if __name__ == "__main__":
    unittest.main()
